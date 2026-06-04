# Day 10 — Full Automation Pipeline
# Author: Baqil Abdelhakim
#
# STEP 1 — Load the raw Excel data
# STEP 2 — Calculate KPIs (counts, averages, filters)
# STEP 3 — Generate charts and save as images
# STEP 4 — Create a new Excel workbook
# STEP 5 — Write data + KPIs into sheets
# STEP 6 — Apply formatting and embed charts

import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image
from datetime import date
# Step 0 - Day 7 
import pandas as pd

data = {
    "supplier":   ["Bosch", "Valeo", "Denso", "Continental", "ZF"],
    "status":     ["Completed", "Pending", "In Progress", "Pending", "Completed"],
    "project":    ["RLVV", "ADAS", "exPHAB", "RLVV", "ADAS"],
    "risk_level": ["Low", "High", "Medium", "High", "Low"],
    "score":      [92, 45, 67, 38, 88]
}
df = pd.DataFrame(data)
df.to_excel("suppliers.xlsx", index=False)

# STEP 1 — Load
df = pd.read_excel("suppliers.xlsx")
print(f"Loaded {len(df)} suppliers")

# STEP 2 — KPIs
total        = len(df)
completed    = len(df[df["status"] == "Completed"])
pending      = len(df[df["status"] == "Pending"])
high_risk    = len(df[df["risk_level"] == "High"])
avg_score    = round(df["score"].mean(), 1)
lowest       = df.loc[df["score"].idxmin(), "supplier"]

print(f"Average score: {avg_score}")
print(f"High risk: {high_risk}")

# STEP 3 — Charts
colors = ["#c0392b" if r == "High" else "#27ae60"
          for r in df["risk_level"]]

plt.figure(figsize=(8, 4))
plt.bar(df["supplier"], df["score"], color=colors)
plt.title("Supplier Scores by Risk Level")
plt.xlabel("Supplier")
plt.ylabel("Score")
plt.ylim(0, 100)
plt.tight_layout()
plt.savefig("chart_scores.png", dpi=150)
plt.close()

status_counts = df["status"].value_counts()
plt.figure(figsize=(5, 5))
plt.pie(status_counts, labels=status_counts.index,
        autopct="%1.0f%%", startangle=90)
plt.title("Status Distribution")
plt.tight_layout()
plt.savefig("chart_status.png", dpi=150)
plt.close()

print("Charts saved.")

# STEP 4 — Create workbook
wb = Workbook()

ws_data = wb.active
ws_data.title = "Raw Data"

ws_kpi = wb.create_sheet("KPI Summary")
ws_dash = wb.create_sheet("Dashboard")

# STEP 5 — Write data
headers = list(df.columns)
ws_data.append(headers)
for _, row in df.iterrows():
    ws_data.append(list(row))

# Write KPIs
kpis = [
    ("Total Suppliers",  total),
    ("Completed",        completed),
    ("Pending",          pending),
    ("High Risk",        high_risk),
    ("Average Score",    avg_score),
    ("Lowest Scorer",    lowest),
    ("Report Date",      date.today().strftime("%d/%m/%Y")),
]
ws_kpi.append(["Metric", "Value"])
for metric, value in kpis:
    ws_kpi.append([metric, value])

print("Data written.")

# STEP 6 — Format Raw Data sheet
green_fill = PatternFill("solid", fgColor="1D9E75")
white_font = Font(bold=True, color="FFFFFF")
red_fill   = PatternFill("solid", fgColor="FFCCCC")

for cell in ws_data[1]:
    cell.fill = green_fill
    cell.font = white_font
    cell.alignment = Alignment(horizontal="center")

for row in ws_data.iter_rows(min_row=2):
    if row[3].value == "High":
        for cell in row:
            cell.fill = red_fill

for col in ws_data.columns:
    max_len = max((len(str(c.value)) for c in col if c.value), default=0)
    ws_data.column_dimensions[col[0].column_letter].width = max_len + 4

ws_data.freeze_panes = "A2"

# Embed charts into Dashboard sheet
img1 = Image("chart_scores.png")
img2 = Image("chart_status.png")
img1.anchor = "B2"
img2.anchor = "B22"
ws_dash.add_image(img1)
ws_dash.add_image(img2)

# Save
wb.save("supplier_full_report.xlsx")
print("Report complete: supplier_full_report.xlsx")